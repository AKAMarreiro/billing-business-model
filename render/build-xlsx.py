"""Generate dist/modelo.xlsx from canonical inputs and derived values."""
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "dados"
DIST = ROOT / "dist"

with (DIST / "derivados.json").open(encoding="utf-8") as handle:
    derived = json.load(handle)
with (DATA / "premissas.json").open(encoding="utf-8") as handle:
    assumptions = json.load(handle)
with (DATA / "base-dispositivos.json").open(encoding="utf-8") as handle:
    base = json.load(handle)

workbook = Workbook()
workbook.remove(workbook.active)
yellow = PatternFill("solid", fgColor="FFF2CC")
header = Font(bold=True, color="FFFFFF")
header_fill = PatternFill("solid", fgColor="0092D6")

def sheet(name, rows, editable=False):
    ws = workbook.create_sheet(name)
    for row_index, row in enumerate(rows, start=1):
        for column_index, value in enumerate(row, start=1):
            cell = ws.cell(row_index, column_index, value)
            if row_index == 1:
                cell.font = header
                cell.fill = header_fill
            elif editable and column_index == 2:
                cell.fill = yellow
    ws.freeze_panes = "A2"
    for column in ws.columns:
        width = max(len(str(cell.value or "")) for cell in column) + 2
        ws.column_dimensions[column[0].column_letter].width = min(width, 42)
    return ws

sheet("Premissas", [["Chave", "Valor"]] + [[key, value] for key, value in assumptions.items()], editable=True)
sheet("Base de Pontos", [["Produto", "Pontos", "Preco anual", "Potencial anual"],
    ["Pivo", base["pivos"]["efetivos"], assumptions["anuidadeRef"], derived["potencialHistorico"]],
    ["Irripump", base["irripump"], assumptions["anuidadeIrripump"], derived["potencialIrripump"]],
    ["Medidor de nivel", base["medidorNivel"], assumptions["anuidadeMedidorNivel"], derived["potencialMedidor"]],
    ["Total", derived["totalEscopo"], derived["precoMedioEscopo"], derived["potencialCompleto"]]])
sheet("Resumo", [["Indicador", "Valor"],
    ["Potencial historico", derived["potencialHistorico"]],
    ["Potencial prospectivo", derived["potencialCompleto"]],
    ["Receita atual", derived["receitaRecorrenteHoje"]],
    ["Gap prospectivo", derived["gapCompleto"]],
    ["Realizacao historica", derived["realizacao2025"]],
    ["Realizacao prospectiva", derived["realizacaoCompleto"]]])
sheet("Stripe Mensal", [["Mes", "Moeda", "Charge", "Invoice"]] + [[row["mes"], row["currency"], row["charge"], row["invoice"]] for row in derived["serieStripe"]])
sheet("Cenarios 5 Anos", [["Cenario", "Ano", "Receita reconhecida", "Margem"]] + [
    [scenario, row["ano"], row["receitaReconhecida"], row["margem"]]
    for scenario, values in derived["cenariosDetalhado"].items() for row in values])

workbook.save(DIST / "modelo.xlsx")
print("OK: dist/modelo.xlsx gerado")
